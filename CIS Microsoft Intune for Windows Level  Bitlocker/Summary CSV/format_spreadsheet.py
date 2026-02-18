import csv
import re
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

INPUT_CSV  = "CIS_LBL_Policy_Summary_FULL.csv"
OUTPUT_XLS = "CIS_LBL_Policy_Summary_FULL.xlsx"

# ── Palette ────────────────────────────────────────────────────────────────────
HEADER_BG      = "1F3864"   # dark navy
HEADER_FG      = "FFFFFF"

ROW_EVEN_BG    = "EEF3FA"   # very light blue-grey
ROW_ODD_BG     = "FFFFFF"

# Complexity badge colours  (background / font)
COMPLEXITY = {
    "Simple":   ("C6EFCE", "276221"),   # green
    "Moderate": ("FFEB9C", "9C6500"),   # amber
    "Complex":  ("FFC7CE", "9C0006"),   # red
}

THIN = Side(style="thin", color="BFC9D6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ── Column widths (characters) ─────────────────────────────────────────────────
COL_WIDTHS = {
    1: 55,   # Name
    2: 70,   # Description
    3: 40,   # Recommended Setting
    4: 13,   # Complexity
    5: 65,   # SMB Justification
}

ILLEGAL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

def clean(value):
    """Strip illegal XML/Excel control characters from a string."""
    if isinstance(value, str):
        return ILLEGAL_CHARS.sub("", value)
    return value

def make_fill(hex_colour):
    return PatternFill("solid", fgColor=hex_colour)

def make_font(bold=False, colour="000000", size=10):
    return Font(bold=bold, color=colour, size=size, name="Segoe UI")

def apply_border(cell):
    cell.border = BORDER

def format_header(cell, text):
    cell.value = text
    cell.font = make_font(bold=True, colour=HEADER_FG, size=10)
    cell.fill = make_fill(HEADER_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    apply_border(cell)

def format_data_cell(cell, row_idx, wrap=True, h_align="left"):
    bg = ROW_EVEN_BG if row_idx % 2 == 0 else ROW_ODD_BG
    cell.fill = make_fill(bg)
    cell.font = make_font(size=9)
    cell.alignment = Alignment(
        horizontal=h_align, vertical="top", wrap_text=wrap
    )
    apply_border(cell)

def format_complexity_cell(cell, value, row_idx):
    bg = ROW_EVEN_BG if row_idx % 2 == 0 else ROW_ODD_BG
    badge = COMPLEXITY.get(value)
    if badge:
        cell.fill = make_fill(badge[0])
        cell.font = make_font(bold=True, colour=badge[1], size=9)
    else:
        cell.fill = make_fill(bg)
        cell.font = make_font(size=9)
    cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=False)
    apply_border(cell)

# ── Main ───────────────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "CIS BL Policies"

with open(INPUT_CSV, encoding="utf-8-sig") as f:
    reader = csv.reader(f)
    rows = list(reader)

headers = rows[0]
data    = rows[1:]

# Header row (row 1)
for col_idx, header in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx)
    format_header(cell, header)

ws.row_dimensions[1].height = 28

# Data rows
for row_idx, row in enumerate(data, start=2):
    # Pad short rows
    while len(row) < len(headers):
        row.append("")

    for col_idx, value in enumerate(row, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=clean(value))

        if col_idx == 4:  # Complexity column
            format_complexity_cell(cell, value.strip(), row_idx)
        else:
            h_align = "left"
            format_data_cell(cell, row_idx, wrap=True, h_align=h_align)

    # Row height: taller rows for long description/justification text
    ws.row_dimensions[row_idx].height = 60

# Column widths
for col_idx, width in COL_WIDTHS.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width

# Freeze the header row
ws.freeze_panes = "A2"

# Auto-filter on header row
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

# ── Sheet tab colour ───────────────────────────────────────────────────────────
ws.sheet_properties.tabColor = "1F3864"

wb.save(OUTPUT_XLS)
print(f"Saved: {OUTPUT_XLS}")
